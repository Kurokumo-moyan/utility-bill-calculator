import csv
import json
import tempfile
import unittest
from datetime import date, datetime, timedelta
from decimal import Decimal as D
from pathlib import Path
from zipfile import ZipFile

from billing import (FIELDS, Source, Tariff, calculate, load_tariffs, money,
                     read_records, sources)

MAPPING = {k: k for k in FIELDS}


class BillingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name)
        self.tariff = Tariff.parse({'usage_rate': '0.30', 'daily_rate': '1', 'data_unit': 'kWh', 'rate_unit': 'kWh', 'interval_minutes': 5})

    def csv(self, name, rows, headers=FIELDS):
        path = self.folder / name
        with path.open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f); writer.writerow(headers); writer.writerows(rows)
        return Source(path)

    def test_full_day_and_duplicate_fixed_fee(self):
        rows = [(datetime(2026, 1, 1) + timedelta(minutes=5*i), '0.1', 'DEMO', 'E1', 'A') for i in range(288)]
        source = self.csv('day.csv', rows)
        result, duplicates = calculate('electricity', [source, source], self.tariff, MAPPING)
        self.assertEqual(duplicates, 288)
        self.assertEqual(result[0].usage, D('28.8'))
        self.assertEqual(result[0].total, D('9.64'))
        self.assertEqual(result[0].status, '完整')

    def test_conflict_rejected(self):
        source = self.csv('conflict.csv', [('2026-01-01', 1, 'D'), ('2026-01-01', 2, 'D')])
        with self.assertRaisesRegex(ValueError, '冲突'):
            calculate('electricity', [source], self.tariff, MAPPING)

    def test_filter_and_fixed_fee_for_missing_days(self):
        source = self.csv('days.csv', [('2026-01-01', 1, 'D'), ('2026-01-03', 2, 'D')])
        result, _ = calculate('electricity', [source], self.tariff, MAPPING, date(2026, 1, 2), date(2026, 1, 4))
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0].total, D('1.6'))
        self.assertIn('不完整', result[0].status)

    def test_water_units_and_gas_daily_only(self):
        source = self.csv('water.csv', [('2026-01-01', '0.1', 'D')])
        water = Tariff.parse({'usage_rate': '.02', 'daily_rate': '.5', 'data_unit': 'kL', 'rate_unit': 'L'})
        result, _ = calculate('water', [source], water, MAPPING)
        self.assertEqual(result[0].total, D('2.5'))
        gas = Tariff.parse({'usage_rate': '0', 'daily_rate': '.6', 'data_unit': 'MJ', 'rate_unit': 'MJ'})
        result, _ = calculate('gas', [source], gas, MAPPING)
        self.assertEqual(result[0].total, D('.6'))

    def test_mapped_headers(self):
        source = self.csv('mapped.csv', [('2026-01-01', '2', 'D')], ['时间', '用量', '表号'])
        mapping = dict(zip(FIELDS[:3], ['时间', '用量', '表号']))
        result, _ = calculate('water', [source], self.tariff, mapping)
        self.assertEqual(result[0].usage, 2)
        with self.assertRaisesRegex(ValueError, '缺少列'):
            calculate('water', [source], self.tariff, MAPPING)

    def test_multiple_meters_rejected(self):
        source = self.csv('meters.csv', [('2026-01-01', 1, 'A'), ('2026-01-01', 2, 'B')])
        with self.assertRaisesRegex(ValueError, '多个表号'):
            calculate('electricity', [source], self.tariff, MAPPING)

    def test_invalid_values(self):
        for value in ('NaN', 'Infinity', '-1', 'not a number', ''):
            with self.subTest(value=value):
                source = self.csv('bad.csv', [('2026-01-01', value, 'D')])
                with self.assertRaises(ValueError):
                    calculate('electricity', [source], self.tariff, MAPPING)

    def test_ambiguous_headers_and_mapping(self):
        source = self.csv('bad-header.csv', [('2026-01-01', 1, 'D', 1)], [*FIELDS[:3], 'ReadValue'])
        with self.assertRaisesRegex(ValueError, '重复'):
            list(read_records(source, MAPPING))
        source = self.csv('mapping.csv', [('2026-01-01', 1, 'D')])
        with self.assertRaisesRegex(ValueError, '同一列'):
            list(read_records(source, {**MAPPING, 'ReadValue': 'ReadDatetime'}))

    def test_tariff_import_and_validation(self):
        path = Path(__file__).resolve().parents[1] / 'examples/rates.example.json'
        currency, tariffs = load_tariffs(path)
        self.assertEqual(currency, 'AUD')
        self.assertEqual(tariffs['water'].daily_rate, D('.50'))
        data = json.loads(path.read_text())
        data['categories']['gas']['daily_rate'] = '-1'
        bad = self.folder / 'bad.json'; bad.write_text(json.dumps(data))
        with self.assertRaises(ValueError):
            load_tariffs(bad)
        for changes in ({'rate_unit': 'L'}, {'interval_minutes': 7}, {'interval_minutes': 2.5}):
            with self.assertRaises(ValueError):
                Tariff.parse({**self.tariff.as_dict(), **changes})

    def test_decimal_rounding(self):
        self.assertEqual(money(D('1.005')), '1.01')
        self.assertEqual(Tariff.parse({**self.tariff.as_dict(), 'rate_unit': 'MJ'}).factor, D('3.6'))

    def test_quality_duplicates_preserve_warning(self):
        source = self.csv('quality.csv', [('2026-01-01', 1, 'D', 'E1', 'A'), ('2026-01-01', 1, 'D', 'E1', 'E')])
        rows, duplicates = calculate('electricity', [source], self.tariff, MAPPING)
        self.assertEqual(duplicates, 1)
        self.assertIn('非A', rows[0].status)

    def test_excel_xlsx_xlsm_multiple_sheets_and_dates(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = 'Day1'
        book.active.append(list(FIELDS))
        book.active.append([datetime(2026, 1, 1), .1, 'DEMO', 'E1', 'A'])
        second = book.create_sheet('Day2')
        second.append(list(FIELDS)); second.append([datetime(2026, 1, 2), .2, 'DEMO', 'E1', 'A'])
        path = self.folder / 'test.xlsx'; book.save(path); book.close()
        macro = self.folder / 'test.xlsm'
        with ZipFile(path) as src, ZipFile(macro, 'w') as dest:
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename == '[Content_Types].xml':
                    data = data.replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml', b'application/vnd.ms-excel.sheet.macroEnabled.main+xml')
                dest.writestr(info, data)
        for file in (path, macro):
            with self.subTest(format=file.suffix):
                selected = sources(file)
                self.assertEqual(len(selected), 2)
                result, _ = calculate('electricity', selected, self.tariff, MAPPING)
                self.assertEqual(sum(r.total for r in result), D('2.09'))
                self.assertEqual(result[0].day, date(2026, 1, 1))

    def test_excel_xls(self):
        import xlwt
        book = xlwt.Workbook(); sheet = book.add_sheet('Usage')
        for col, value in enumerate(FIELDS):
            sheet.write(0, col, value)
        sheet.write(1, 0, datetime(2026, 1, 1), xlwt.easyxf(num_format_str='YYYY-MM-DD HH:MM:SS'))
        for col, value in enumerate([.1, 'DEMO', 'E1', 'A'], 1):
            sheet.write(1, col, value)
        path = self.folder / 'test.xls'; book.save(str(path))
        rows, _ = calculate('electricity', sources(path), self.tariff, MAPPING)
        self.assertEqual(rows[0].total, D('1.03'))


if __name__ == '__main__':
    unittest.main()
